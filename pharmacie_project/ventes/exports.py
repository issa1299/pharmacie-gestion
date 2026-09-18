"""Export Excel (.xlsx) des listes — avec mise en forme professionnelle."""
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse

from .models import Vente
from medicaments.models import Medicament
from clients.models import Client
from stock.models import MouvementStock


def _workbook():
    """Creer un classeur Excel avec format commun."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    return wb, {
        'header_font': Font(name='Calibri', bold=True, color='FFFFFF', size=11),
        'header_fill': PatternFill(start_color='0F6E56', end_color='0F6E56', fill_type='solid'),
        'header_align': Alignment(horizontal='center', vertical='center', wrap_text=True),
        'cell_font': Font(name='Calibri', size=10),
        'cell_align': Alignment(vertical='center'),
        'num_align': Alignment(horizontal='right', vertical='center'),
        'date_align': Alignment(horizontal='center', vertical='center'),
        'border': Border(
            left=Side(style='thin', color='D0D0D0'),
            right=Side(style='thin', color='D0D0D0'),
            top=Side(style='thin', color='D0D0D0'),
            bottom=Side(style='thin', color='D0D0D0'),
        ),
        'alt_fill': PatternFill(start_color='F4F8F6', end_color='F4F8F6', fill_type='solid'),
        'green_font': Font(name='Calibri', size=10, color='0F6E56', bold=True),
        'red_font': Font(name='Calibri', size=10, color='C0392B', bold=True),
        'orange_font': Font(name='Calibri', size=10, color='E65100', bold=True),
    }


def _format_sheet(ws, headers, widths, fmt):
    """Appliquer le formatage aux en-tetes et aux colonnes."""
    # En-tetes
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = fmt['header_font']
        cell.fill = fmt['header_fill']
        cell.alignment = fmt['header_align']
        cell.border = fmt['border']

    # Largeurs de colonnes
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else 'A' + chr(64 + col_idx - 26)].width = width

    # Figer la ligne 1
    ws.freeze_panes = 'A2'

    # Filtre auto
    ws.auto_filter.ref = ws.dimensions


def _apply_row_style(ws, row_idx, num_cols, fmt, is_alt=False):
    """Appliser le style a une ligne de donnees."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = fmt['cell_font']
        cell.alignment = fmt['cell_align']
        cell.border = fmt['border']
        if is_alt:
            cell.fill = fmt['alt_fill']


@login_required
def export_ventes_xlsx(request):
    """Export des ventes (1 ligne par vente)."""
    wb, fmt = _workbook()
    ws = wb.active
    ws.title = 'Ventes'

    headers = [
        'N Facture', 'Date', 'Client', 'Caissier', 'Articles',
        'Total', 'Remise', 'Net a payer', 'Mode paiement',
        'Statut paiement', 'Statut',
    ]
    widths = [18, 18, 22, 20, 10, 14, 14, 14, 16, 16, 12]
    _format_sheet(ws, headers, widths, fmt)

    ventes = Vente.objects.select_related('client', 'utilisateur').order_by('-date_vente')
    for row_idx, v in enumerate(ventes, 2):
        nb_articles = sum(l.quantite for l in v.lignes.all())
        ws.cell(row=row_idx, column=1, value=v.numero_facture)
        ws.cell(row=row_idx, column=2, value=v.date_vente.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row_idx, column=3, value=str(v.client) if v.client else '—')
        ws.cell(row=row_idx, column=4, value=v.utilisateur.get_full_name() or v.utilisateur.username if v.utilisateur else '—')
        ws.cell(row=row_idx, column=5, value=nb_articles)
        ws.cell(row=row_idx, column=6, value=float(v.total))
        ws.cell(row=row_idx, column=7, value=float(v.remise))
        ws.cell(row=row_idx, column=8, value=float(v.net_a_payer))
        ws.cell(row=row_idx, column=9, value=v.get_mode_paiement_display())
        ws.cell(row=row_idx, column=10, value=v.get_statut_paiement_display())
        statut_cell = ws.cell(row=row_idx, column=11, value=v.get_statut_display())
        if v.statut == 'validee':
            statut_cell.font = fmt['green_font']
        else:
            statut_cell.font = fmt['red_font']
        # Format monetaire
        for col in [6, 7, 8]:
            ws.cell(row=row_idx, column=col).number_format = '#,##0'
            ws.cell(row=row_idx, column=col).alignment = fmt['num_align']
        ws.cell(row=row_idx, column=5).alignment = fmt['num_align']
        _apply_row_style(ws, row_idx, 11, fmt, row_idx % 2 == 0)

    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    resp['Content-Disposition'] = 'attachment; filename="ventes.xlsx"'
    wb.save(resp)
    return resp


@login_required
def export_vente_detail_xlsx(request, pk):
    """Export du detail d'une vente."""
    from django.shortcuts import get_object_or_404
    vente = get_object_or_404(Vente, pk=pk)
    wb, fmt = _workbook()
    ws = wb.active
    ws.title = 'Detail vente'

    headers = ['N Facture', 'Date', 'Medicament', 'Code-barres', 'Quantite', 'Prix unitaire', 'Sous-total']
    widths = [18, 18, 28, 18, 10, 14, 14]
    _format_sheet(ws, headers, widths, fmt)

    for row_idx, l in enumerate(vente.lignes.select_related('medicament').all(), 2):
        ws.cell(row=row_idx, column=1, value=vente.numero_facture)
        ws.cell(row=row_idx, column=2, value=vente.date_vente.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row_idx, column=3, value=l.medicament.nom)
        ws.cell(row=row_idx, column=4, value=l.medicament.code_barre or '—')
        ws.cell(row=row_idx, column=5, value=l.quantite)
        ws.cell(row=row_idx, column=6, value=float(l.prix_unitaire))
        ws.cell(row=row_idx, column=7, value=float(l.sous_total))
        for col in [5, 6, 7]:
            ws.cell(row=row_idx, column=col).number_format = '#,##0'
            ws.cell(row=row_idx, column=col).alignment = fmt['num_align']
        _apply_row_style(ws, row_idx, 7, fmt, row_idx % 2 == 0)

    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    resp['Content-Disposition'] = f'attachment; filename="{vente.numero_facture}.xlsx"'
    wb.save(resp)
    return resp


@login_required
def export_medicaments_xlsx(request):
    """Export du catalogue complet."""
    wb, fmt = _workbook()
    ws = wb.active
    ws.title = 'Medicaments'

    headers = [
        'Nom', 'Categorie', 'Etagere', 'Code-barres', 'Description',
        'Prix achat', 'Prix vente', 'Marge',
        'Stock', 'Seuil alerte', 'Date expiration', 'Statut',
    ]
    widths = [28, 20, 18, 18, 30, 14, 14, 14, 10, 12, 16, 14]
    _format_sheet(ws, headers, widths, fmt)

    from django.utils import timezone
    today = timezone.now().date()
    meds = Medicament.objects.select_related('categorie', 'etagere').order_by('nom')

    for row_idx, m in enumerate(meds, 2):
        if m.date_expiration and m.date_expiration < today:
            statut = 'Expire'
            statut_font = fmt['red_font']
        elif m.quantite_stock <= m.seuil_alerte:
            statut = 'Stock faible'
            statut_font = fmt['orange_font']
        else:
            statut = 'OK'
            statut_font = fmt['green_font']

        marge = float(m.prix_vente - m.prix_achat) if m.prix_achat else 0

        ws.cell(row=row_idx, column=1, value=m.nom)
        ws.cell(row=row_idx, column=2, value=m.categorie.nom if m.categorie else '—')
        ws.cell(row=row_idx, column=3, value=f"{m.etagere.code} — {m.etagere.nom}" if m.etagere else '—')
        ws.cell(row=row_idx, column=4, value=m.code_barre or '—')
        ws.cell(row=row_idx, column=5, value=m.description or '—')
        ws.cell(row=row_idx, column=6, value=float(m.prix_achat))
        ws.cell(row=row_idx, column=7, value=float(m.prix_vente))
        ws.cell(row=row_idx, column=8, value=marge)
        ws.cell(row=row_idx, column=9, value=m.quantite_stock)
        ws.cell(row=row_idx, column=10, value=m.seuil_alerte)
        ws.cell(row=row_idx, column=11, value=m.date_expiration.strftime('%d/%m/%Y') if m.date_expiration else '—')
        statut_cell = ws.cell(row=row_idx, column=12, value=statut)
        statut_cell.font = statut_font

        for col in [6, 7, 8]:
            ws.cell(row=row_idx, column=col).number_format = '#,##0'
            ws.cell(row=row_idx, column=col).alignment = fmt['num_align']
        for col in [9, 10]:
            ws.cell(row=row_idx, column=col).alignment = fmt['num_align']
        _apply_row_style(ws, row_idx, 12, fmt, row_idx % 2 == 0)

    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    resp['Content-Disposition'] = 'attachment; filename="medicaments.xlsx"'
    wb.save(resp)
    return resp


@login_required
def export_clients_xlsx(request):
    """Export des clients avec statistiques d'achat."""
    from django.db.models import Count, Q, Sum
    wb, fmt = _workbook()
    ws = wb.active
    ws.title = 'Clients'

    headers = ['Nom', 'Prenom', 'Telephone', 'Email', 'Adresse', 'Date naissance', 'Nb achats', 'Total depense']
    widths = [20, 18, 16, 28, 28, 16, 12, 16]
    _format_sheet(ws, headers, widths, fmt)

    ventes_validees = Q(ventes__statut='validee')
    clients = Client.objects.annotate(
        nb_ventes=Count('ventes', filter=ventes_validees),
        _total_brut=Sum('ventes__total', filter=ventes_validees),
        _total_remises=Sum('ventes__remise', filter=ventes_validees),
    ).order_by('nom')

    for row_idx, c in enumerate(clients, 2):
        total = float((c._total_brut or 0) - (c._total_remises or 0))
        ws.cell(row=row_idx, column=1, value=c.nom)
        ws.cell(row=row_idx, column=2, value=c.prenom)
        ws.cell(row=row_idx, column=3, value=c.telephone or '—')
        ws.cell(row=row_idx, column=4, value=c.email or '—')
        ws.cell(row=row_idx, column=5, value=c.adresse or '—')
        ws.cell(row=row_idx, column=6, value=c.date_naissance.strftime('%d/%m/%Y') if c.date_naissance else '—')
        ws.cell(row=row_idx, column=7, value=c.nb_ventes)
        ws.cell(row=row_idx, column=8, value=total)
        ws.cell(row=row_idx, column=7).alignment = fmt['num_align']
        ws.cell(row=row_idx, column=8).number_format = '#,##0'
        ws.cell(row=row_idx, column=8).alignment = fmt['num_align']
        _apply_row_style(ws, row_idx, 8, fmt, row_idx % 2 == 0)

    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    resp['Content-Disposition'] = 'attachment; filename="clients.xlsx"'
    wb.save(resp)
    return resp


@login_required
def export_mouvements_xlsx(request):
    """Export du journal de stock."""
    wb, fmt = _workbook()
    ws = wb.active
    ws.title = 'Mouvements'

    headers = ['Date', 'Medicament', 'Type', 'Quantite', 'Fournisseur', 'Motif', 'Utilisateur']
    widths = [18, 28, 14, 10, 22, 24, 20]
    _format_sheet(ws, headers, widths, fmt)

    for row_idx, m in enumerate(
        MouvementStock.objects.select_related('medicament', 'utilisateur', 'fournisseur').order_by('-date'),
        2
    ):
        ws.cell(row=row_idx, column=1, value=m.date.strftime('%d/%m/%Y %H:%M'))
        ws.cell(row=row_idx, column=2, value=m.medicament.nom)
        type_cell = ws.cell(row=row_idx, column=3, value=m.get_type_mouvement_display())
        if m.type_mouvement == 'entree':
            type_cell.font = fmt['green_font']
        elif m.type_mouvement == 'sortie':
            type_cell.font = fmt['red_font']
        ws.cell(row=row_idx, column=4, value=m.quantite)
        ws.cell(row=row_idx, column=4).alignment = fmt['num_align']
        ws.cell(row=row_idx, column=5, value=m.fournisseur.nom if m.fournisseur else '—')
        ws.cell(row=row_idx, column=6, value=m.motif or '—')
        ws.cell(row=row_idx, column=7, value=m.utilisateur.get_full_name() or m.utilisateur.username if m.utilisateur else '—')
        _apply_row_style(ws, row_idx, 7, fmt, row_idx % 2 == 0)

    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    resp['Content-Disposition'] = 'attachment; filename="mouvements_stock.xlsx"'
    wb.save(resp)
    return resp
